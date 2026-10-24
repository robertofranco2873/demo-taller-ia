from io import BytesIO

import pytest
from openpyxl import load_workbook

from src.export import quote_to_excel
from src.quote import Quote, QuoteLine


@pytest.fixture
def quote():
    lines = [
        QuoteLine("S01", "Cambio de aceite", std_hours=0.5, parts_cost=880, qty=1),
        QuoteLine("S03", "Balatas", std_hours=1.5, parts_cost=950, qty=2),
    ]
    return Quote(lines, labor_rate=400, margin_pct=0.10)


def test_quote_formula(quote):
    # Mano de obra = (0.5 + 1.5×2) h × 400 = 1,400
    assert quote.hours == pytest.approx(3.5)
    assert quote.labor == pytest.approx(1400)
    # Refacciones = 880 + 950×2 = 2,780
    assert quote.parts == pytest.approx(2780)
    assert quote.base == pytest.approx(4180)
    assert quote.margin == pytest.approx(418)
    assert quote.subtotal == pytest.approx(4598)
    assert quote.iva == pytest.approx(735.68)
    assert quote.total == pytest.approx(5333.68)


def test_line_subtotals_add_up_to_subtotal(quote):
    assert sum(quote.line_subtotal(l) for l in quote.lines) == pytest.approx(quote.subtotal)


def test_empty_quote_is_zero():
    q = Quote([], labor_rate=450, margin_pct=0.15)
    assert q.total == 0


@pytest.mark.parametrize("rate, margin", [(-1, 0.1), (100, -0.1)])
def test_invalid_parameters(rate, margin):
    with pytest.raises(ValueError):
        Quote([], labor_rate=rate, margin_pct=margin)


def test_invalid_quantity():
    with pytest.raises(ValueError):
        Quote([QuoteLine("S01", "x", 1, 0, qty=0)], 100, 0)


def test_excel_export_contains_total(quote):
    wb = load_workbook(BytesIO(quote_to_excel(quote, "Cliente X", "Auto Y")))
    ws = wb.active
    values = [c.value for row in ws.iter_rows() for c in row]
    assert "TOTAL" in values
    assert quote.total in values
    assert any(isinstance(v, str) and "ficticios" in v for v in values)
