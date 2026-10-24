"""Exportación de la cotización a Excel (openpyxl)."""
from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from src import clock
from src.config import BUSINESS_NAME, DEMO_NOTICE
from src.quote import Quote

MONEY = '"$"#,##0.00'


def quote_to_excel(quote: Quote, customer: str = "", vehicle: str = "",
                   when: datetime | None = None) -> bytes:
    when = when or clock.now()
    wb = Workbook()
    ws = wb.active
    ws.title = "Cotización"
    ws["A1"] = f"{BUSINESS_NAME} — Cotización"
    ws["A1"].font = Font(size=14, bold=True)
    ws["A2"] = DEMO_NOTICE
    ws["A2"].font = Font(italic=True, color="C00000")
    ws["A4"], ws["B4"] = "Fecha", when.strftime("%Y-%m-%d %H:%M")
    ws["A5"], ws["B5"] = "Cliente", customer or "—"
    ws["A6"], ws["B6"] = "Vehículo", vehicle or "—"
    ws["A7"], ws["B7"] = "Tarifa por hora", quote.labor_rate
    ws["B7"].number_format = MONEY
    ws["A8"], ws["B8"] = "Margen", quote.margin_pct
    ws["B8"].number_format = "0.0%"

    header = ["Código", "Servicio", "Cantidad", "Horas estándar", "Mano de obra", "Refacciones", "Importe (con margen)"]
    ws.append([])
    ws.append(header)
    head_row = ws.max_row
    for cell in ws[head_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center")
    for line in quote.lines:
        ws.append([line.code, line.name, line.qty, round(line.hours, 2),
                   round(line.labor(quote.labor_rate), 2), round(line.parts(), 2),
                   quote.line_subtotal(line)])
        for col in "EFG":
            ws[f"{col}{ws.max_row}"].number_format = MONEY

    ws.append([])
    for label, value in [("Mano de obra", quote.labor), ("Refacciones", quote.parts),
                         ("Margen", quote.margin), ("Subtotal", quote.subtotal),
                         (f"IVA {quote.iva_rate:.0%}", quote.iva), ("TOTAL", quote.total)]:
        ws.append(["", "", "", "", "", label, value])
        ws[f"G{ws.max_row}"].number_format = MONEY
        if label == "TOTAL":
            ws[f"F{ws.max_row}"].font = ws[f"G{ws.max_row}"].font = Font(bold=True)
    ws.append([])
    ws.append(["Fórmula: horas estándar × tarifa + refacciones; + margen; + IVA. Vigencia: 7 días."])

    for col, width in zip("ABCDEFG", [10, 38, 10, 15, 15, 15, 20]):
        ws.column_dimensions[col].width = width
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _latin1(text: str) -> str:
    """Las fuentes base del PDF son latin-1: se reemplazan símbolos fuera de ese juego."""
    text = text.replace("—", "-").replace("×", "x").replace("⚠️", "").replace("⚠", "")
    return text.encode("latin-1", "replace").decode("latin-1").strip()


def quote_to_pdf(quote: Quote, customer: str = "", vehicle: str = "",
                 when: datetime | None = None) -> bytes:
    from fpdf import FPDF

    when = when or clock.now()
    pdf = FPDF(format="Letter")
    pdf.set_auto_page_break(True, margin=15)
    pdf.add_page()
    pdf.set_fill_color(122, 100, 180)
    pdf.rect(0, 0, pdf.w, 6, "F")
    pdf.set_font("Helvetica", "B", 18)
    pdf.set_text_color(40, 36, 60)
    pdf.cell(0, 12, _latin1(f"{BUSINESS_NAME} - Cotización"), new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "I", 9)
    pdf.set_text_color(170, 40, 40)
    pdf.multi_cell(0, 5, _latin1(DEMO_NOTICE), new_x="LMARGIN", new_y="NEXT")
    pdf.ln(2)
    pdf.set_text_color(40, 36, 60)
    pdf.set_font("Helvetica", "", 10)
    for label, value in [("Fecha", when.strftime("%Y-%m-%d %H:%M")), ("Cliente", customer or "-"),
                         ("Vehículo", vehicle or "-"), ("Tarifa por hora", f"${quote.labor_rate:,.2f}"),
                         ("Margen", f"{quote.margin_pct:.1%}")]:
        pdf.cell(35, 6, _latin1(label), new_x="RIGHT", new_y="TOP")
        pdf.cell(0, 6, _latin1(value), new_x="LMARGIN", new_y="NEXT")
    pdf.ln(3)

    widths = [16, 70, 14, 22, 26, 26, 22]
    header = ["Código", "Servicio", "Cant.", "Horas est.", "Mano de obra", "Refacciones", "Importe"]
    pdf.set_font("Helvetica", "B", 9)
    pdf.set_fill_color(62, 92, 140)
    pdf.set_text_color(255, 255, 255)
    for w, h in zip(widths, header):
        pdf.cell(w, 7, _latin1(h), border=0, align="C", fill=True)
    pdf.ln()
    pdf.set_font("Helvetica", "", 9)
    pdf.set_text_color(40, 36, 60)
    for i, line in enumerate(quote.lines):
        pdf.set_fill_color(244, 241, 250) if i % 2 == 0 else pdf.set_fill_color(255, 255, 255)
        cells = [line.code, line.name[:42], str(line.qty), f"{line.hours:.2f}",
                 f"${line.labor(quote.labor_rate):,.2f}", f"${line.parts():,.2f}",
                 f"${quote.line_subtotal(line):,.2f}"]
        for w, c, align in zip(widths, cells, "LLCRRRR"):
            pdf.cell(w, 7, _latin1(c), align=align, fill=True)
        pdf.ln()
    pdf.ln(3)
    for label, value in [("Mano de obra", quote.labor), ("Refacciones", quote.parts),
                         ("Margen", quote.margin), ("Subtotal", quote.subtotal),
                         (f"IVA {quote.iva_rate:.0%}", quote.iva), ("TOTAL", quote.total)]:
        pdf.set_font("Helvetica", "B" if label == "TOTAL" else "", 11 if label == "TOTAL" else 10)
        pdf.cell(sum(widths) - 40, 6, _latin1(label), align="R")
        pdf.cell(40, 6, f"${value:,.2f}", align="R", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(4)
    pdf.set_font("Helvetica", "I", 8)
    pdf.multi_cell(0, 4, _latin1("Fórmula: (horas estándar x tarifa + refacciones) x (1 + margen) x (1 + IVA). "
                                 "Vigencia: 7 días. Documento de demostración con datos ficticios."))
    return bytes(pdf.output())
